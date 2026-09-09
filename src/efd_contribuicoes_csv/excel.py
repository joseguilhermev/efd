from __future__ import annotations

import csv
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SHEETS = (
    ("Analítico", "efd_contribuicoes_analitico.csv"),
    ("Indicadores", "efd_contribuicoes_indicadores.csv"),
    ("Comparação", "efd_comparacao_notas.csv"),
    ("Não lançadas", "efd_icms_nao_lancadas_contribuicoes.csv"),
    ("Pendências", "efd_pendencias_conferencia.csv"),
    ("Cobertura", "efd_cobertura_registros.csv"),
    ("Períodos", "efd_periodos_escopo.csv"),
)

MONEY_COLUMNS = {
    "Valor Operação",
    "Valor PIS",
    "Valor Cofins",
    "Valor Documento EFD Contribuições",
    "Valor Documento EFD ICMS",
}
INTEGER_COLUMNS = {
    "Quantidade Registros",
    "Quantidade EFD Contribuições",
    "Quantidade EFD ICMS",
    "Primeira Linha",
}
DATE_COLUMN_PREFIXES = ("Data Documento", "Data Entrada/Saída")


def _is_money_column(name: str) -> bool:
    return name.startswith("Vlr ") or name in MONEY_COLUMNS


def _is_decimal_column(name: str) -> bool:
    return name == "Qtde" or name.startswith("Qtde ") or name.startswith("Alíquota ")


def _decimal(value: str) -> Decimal:
    normalized = value.replace(".", "").replace(",", ".") if "," in value else value
    return Decimal(normalized)


def _typed_value(column: str, value: str) -> object:
    # Remove apenas controles rejeitados pelo Excel; o CSV original é preservado.
    value = ILLEGAL_CHARACTERS_RE.sub("", value).strip()
    if not value:
        return None
    if column.startswith(DATE_COLUMN_PREFIXES):
        for date_format in ("%d/%m/%Y", "%d%m%Y"):
            try:
                return datetime.strptime(value, date_format).date()
            except ValueError:
                continue
        return value
    if column in INTEGER_COLUMNS:
        try:
            return int(value)
        except ValueError:
            return value
    if _is_money_column(column) or _is_decimal_column(column):
        try:
            return _decimal(value)
        except InvalidOperation:
            return value
    return value


def _number_format(column: str) -> str:
    if column.startswith(DATE_COLUMN_PREFIXES):
        return "dd/mm/yyyy"
    if column in INTEGER_COLUMNS:
        return "#,##0"
    if _is_money_column(column):
        return "#,##0.00"
    if _is_decimal_column(column):
        return "#,##0.0000"
    return "@"


def create_excel_workbook(
    output_directory: str | Path,
    workbook_path: str | Path,
    *,
    delimiter: str = ";",
) -> Path:
    """Cria uma pasta de trabalho tipada a partir das saídas CSV."""

    output = Path(output_directory)
    destination = Path(workbook_path)
    workbook = Workbook()
    workbook.remove(workbook.active)

    for sheet_name, filename in SHEETS:
        worksheet = workbook.create_sheet(sheet_name)
        with (output / filename).open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.reader(stream, delimiter=delimiter)
            rows = iter(reader)
            headers = next(rows)
            worksheet.append(headers)
            for row in rows:
                worksheet.append(
                    [_typed_value(column, value) for column, value in zip(headers, row)]
                )

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.sheet_view.showGridLines = False
        for cell in worksheet[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.row_dimensions[1].height = 30

        for index, header in enumerate(headers, start=1):
            letter = get_column_letter(index)
            for cell in worksheet[letter][1:]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.number_format = _number_format(header)
            longest = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in worksheet[letter]
            )
            worksheet.column_dimensions[letter].width = min(max(longest + 2, 11), 35)

    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return destination
