from __future__ import annotations

import csv
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from efd_contribuicoes_csv.workflow import (
    SCOPE_COLUMNS,
    EFDInventoryError,
    discover_annual_efd_input,
    process_annual_efd_input,
    process_efd_files,
)
from efd_contribuicoes_csv.workflow_cli import main

PROJECT_ROOT = Path(__file__).parents[1]
CONTRIBUTION = PROJECT_ROOT / "efd_contribuicoes_sintetico_1000_linhas.txt"
ICMS = PROJECT_ROOT / "efd_icms_sintetico_estrutura_real.txt"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream, delimiter=";"))


def create_annual_input(path: Path) -> Path:
    contribution_directory = path / "efd_contribuicoes"
    icms_directory = path / "efd_icms"
    contribution_directory.mkdir(parents=True)
    icms_directory.mkdir()
    shutil.copyfile(CONTRIBUTION, contribution_directory / "efd_08_2026.txt")
    shutil.copyfile(ICMS, icms_directory / "efd_08_2026.txt")
    return path


def test_runs_full_flow_and_marks_missing_scope_periods(tmp_path: Path) -> None:
    output = tmp_path / "result"

    result = process_efd_files(
        CONTRIBUTION,
        ICMS,
        output,
        period_start="08/2026",
        period_end="10/2026",
    )
    scope_rows = read_csv(result.scope_path)
    missing_notes = read_csv(result.missing_notes_path)
    indicator_rows = read_csv(result.indicators.output_path)
    zero_rows = [row for row in indicator_rows if row["Quantidade Registros"] == "0"]

    assert result.conversion.rows == 611
    assert result.comparison.rows == 221
    assert result.missing_notes == 0
    assert result.workbook_path.is_file()
    assert missing_notes == []
    assert result.pending_notes == 1
    assert "F550" in read_csv(result.pending_notes_path)[0]["Evidências EFD Contribuições"]
    assert tuple(scope_rows[0]) == SCOPE_COLUMNS
    assert [row["EFD Contribuições"] for row in scope_rows] == [
        "PROCESSADO",
        "AUSENTE",
        "AUSENTE",
    ]
    assert len(zero_rows) == 9
    assert {row["Período"] for row in zero_rows} == {
        "01/08/2026 a 31/08/2026",
        "01/09/2026 a 30/09/2026",
        "01/10/2026 a 31/10/2026",
    }

    workbook = load_workbook(result.workbook_path, read_only=True, data_only=True)
    assert workbook.sheetnames == [
        "Analítico",
        "Indicadores",
        "Comparação",
        "Não lançadas",
        "Pendências",
        "Cobertura",
        "Períodos",
    ]
    analytical = workbook["Analítico"]
    headers = {cell.value: cell.column for cell in analytical[1]}
    assert analytical.cell(2, headers["CNPJ"]).value == "99999999000199"
    assert analytical.cell(2, headers["CNPJ"]).data_type == "s"
    assert analytical.cell(2, headers["Chave NF-e"]).data_type == "s"
    assert analytical.cell(2, headers["Vlr Documento"]).data_type == "n"
    assert analytical.cell(2, headers["Data Documento"]).data_type == "d"
    comparison = workbook["Comparação"]
    comparison_headers = {cell.value: cell.column for cell in comparison[1]}
    comparison_key = next(
        row[comparison_headers["Chave NF-e"] - 1]
        for row in comparison.iter_rows(min_row=2)
        if row[comparison_headers["Chave NF-e"] - 1].value == "35260899000000010001553230000200231700200236"
    )
    assert comparison_key.data_type == "s"


def test_discovers_and_processes_available_periods_for_the_whole_year(
    tmp_path: Path,
) -> None:
    input_directory = create_annual_input(tmp_path / "entrada")
    output = tmp_path / "resultado"

    inventory = discover_annual_efd_input(input_directory)
    result = process_annual_efd_input(inventory, output)
    scope_rows = read_csv(result.scope_path)
    indicator_rows = read_csv(result.indicators.output_path)

    assert inventory.year == 2026
    assert len(inventory.contributions) == 1
    assert len(inventory.icms) == 1
    assert len(inventory.missing_contributions) == 11
    assert len(inventory.missing_icms) == 11
    assert result.conversion.rows == 611
    assert result.comparison.rows == 221
    assert len(scope_rows) == 12
    assert scope_rows[7]["Status"] == "PROCESSADO"
    assert scope_rows[0]["EFD Contribuições"] == "AUSENTE"
    assert {row["Período"] for row in indicator_rows} == {
        row["Período"] for row in scope_rows
    }


def test_accepts_matrix_and_branch_with_the_same_cnpj_root(tmp_path: Path) -> None:
    input_directory = create_annual_input(tmp_path / "entrada")
    icms_path = input_directory / "efd_icms" / "efd_08_2026.txt"
    content = icms_path.read_text(encoding="utf-8").replace(
        "99999999000199", "99999999000270", 1
    )
    icms_path.write_text(content, encoding="utf-8")

    inventory = discover_annual_efd_input(input_directory)

    assert inventory.cnpj == "99999999000199"
    assert len(inventory.contributions) == 1
    assert len(inventory.icms) == 1


def test_main_reports_missing_months_and_cancels_before_writing(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    input_directory = create_annual_input(tmp_path / "entrada")
    output = tmp_path / "resultado"
    monkeypatch.setattr("builtins.input", lambda _: "n")

    exit_code = main([str(input_directory), "-o", str(output)])
    message = capsys.readouterr().out

    assert exit_code == 1
    assert "EFD Contribuições ausente: 01/2026" in message
    assert "EFD ICMS/IPI ausente: 01/2026" in message
    assert "Processamento cancelado" in message
    assert not output.exists()


def test_main_can_continue_non_interactively_with_missing_months(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    input_directory = create_annual_input(tmp_path / "entrada")
    output = tmp_path / "resultado"

    exit_code = main(
        [
            str(input_directory),
            "-o",
            str(output),
            "--continuar-com-ausentes",
        ]
    )
    message = capsys.readouterr().out

    assert exit_code == 0
    assert "Fluxo concluído" in message
    assert (output / "efd_contribuicoes_analitico.csv").is_file()
    assert (output / "efd_contribuicoes_indicadores.csv").is_file()
    assert (output / "efd_comparacao_notas.csv").is_file()
    assert (output / "efd_periodos_escopo.csv").is_file()
    assert (output / "efd_resultado.xlsx").is_file()


def test_processes_available_data_when_months_do_not_have_a_pair(
    tmp_path: Path,
) -> None:
    input_directory = create_annual_input(tmp_path / "entrada")
    icms_path = input_directory / "efd_icms" / "efd_08_2026.txt"
    september = icms_path.read_text(encoding="utf-8")
    september = september.replace("01082026", "01092026", 1)
    september = september.replace("31082026", "30092026", 1)
    icms_path.write_text(september, encoding="utf-8")
    output = tmp_path / "resultado"

    inventory = discover_annual_efd_input(input_directory)
    result = process_annual_efd_input(inventory, output)
    scope_rows = read_csv(result.scope_path)

    assert result.conversion.rows == 611
    assert result.comparison.rows == 221
    assert result.comparison.by_status["SEM_EFD_CONTRIBUICOES"] == 1
    assert result.comparison.by_status["REVISAO_NECESSARIA"] == 118
    assert result.missing_notes == 0
    assert result.pending_notes == 221
    assert scope_rows[7]["EFD Contribuições"] == "PROCESSADO"
    assert scope_rows[7]["EFD ICMS/IPI"] == "AUSENTE"
    assert scope_rows[8]["EFD Contribuições"] == "AUSENTE"
    assert scope_rows[8]["EFD ICMS/IPI"] == "PROCESSADO"


def test_reports_year_cnpj_and_duplicate_period_inconsistencies(
    tmp_path: Path,
) -> None:
    input_directory = create_annual_input(tmp_path / "entrada")
    contribution_directory = input_directory / "efd_contribuicoes"
    icms_directory = input_directory / "efd_icms"
    duplicate = CONTRIBUTION.read_text(encoding="utf-8").replace(
        "99999999000199", "11111111000111", 1
    )
    (contribution_directory / "duplicada_08_2026.txt").write_text(
        duplicate, encoding="utf-8"
    )
    different_year = ICMS.read_text(encoding="utf-8").replace("2026", "2025")
    (icms_directory / "efd_08_2025.txt").write_text(
        different_year, encoding="utf-8"
    )

    with pytest.raises(EFDInventoryError) as captured:
        discover_annual_efd_input(input_directory)

    message = str(captured.value)
    assert "mais de um ano" in message
    assert "CNPJs de raízes diferentes" in message
    assert "duplicada em 08/2026" in message


def test_missing_notes_include_duplicates_exclusive_to_icms(tmp_path: Path) -> None:
    source = tmp_path / "icms.txt"
    lines = ICMS.read_text(encoding="utf-8").splitlines(keepends=True)
    last_note = next(line for line in reversed(lines) if line.startswith("|C100|"))
    last_note_index = max(index for index, line in enumerate(lines) if line.startswith("|C100|"))
    lines.insert(last_note_index, last_note)
    source.write_text("".join(lines), encoding="utf-8")
    contribution = tmp_path / "contribution.txt"
    contribution.write_text("\n".join(
        line for line in CONTRIBUTION.read_text(encoding="utf-8").splitlines()
        if not line.startswith("|F550|")
    ), encoding="utf-8")
    result = process_efd_files(contribution, source, tmp_path / "resultado")
    missing = read_csv(result.missing_notes_path)
    assert result.missing_notes == 1
    assert missing[0]["Status"] == "DUPLICADA_EFD_ICMS"
    assert missing[0]["Quantidade EFD Contribuições"] == "0"
    assert missing[0]["Quantidade EFD ICMS"] == "2"


@pytest.mark.parametrize("annual", [False, True])
def test_invalid_amount_preserves_previous_outputs(tmp_path: Path, annual: bool) -> None:
    input_directory = create_annual_input(tmp_path / "entrada")
    contribution = input_directory / "efd_contribuicoes" / "efd_08_2026.txt"
    icms = input_directory / "efd_icms" / "efd_08_2026.txt"
    lines = contribution.read_text(encoding="utf-8").splitlines()
    for index, line in enumerate(lines):
        if line.startswith("|A170|"):
            fields = line.split("|")
            fields[5] = "valor-inválido"
            lines[index] = "|".join(fields)
    contribution.write_text("\n".join(lines), encoding="utf-8")
    output = tmp_path / "resultado"
    output.mkdir()
    previous = output / "efd_contribuicoes_analitico.csv"
    previous.write_text("resultado anterior", encoding="utf-8")
    with pytest.raises(ValueError, match="valor inválido"):
        if annual:
            process_annual_efd_input(discover_annual_efd_input(input_directory), output)
        else:
            process_efd_files(contribution, icms, output)
    assert previous.read_text(encoding="utf-8") == "resultado anterior"
    assert list(output.iterdir()) == [previous]


def test_workflow_does_not_overwrite_input_with_report(tmp_path: Path) -> None:
    source = tmp_path / "efd_contribuicoes_analitico.csv"
    shutil.copyfile(CONTRIBUTION, source)
    original = source.read_bytes()
    with pytest.raises(ValueError, match="sobrescrever um arquivo EFD"):
        process_efd_files(source, ICMS, tmp_path)
    assert source.read_bytes() == original
    assert list(tmp_path.iterdir()) == [source]
